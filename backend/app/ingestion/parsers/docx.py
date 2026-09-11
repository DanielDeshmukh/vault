import base64
from io import BytesIO


class DOCXParser:
    """Parse Word documents."""
    
    async def parse(self, content: str) -> str:
        """
        Parse DOCX content.
        
        Args:
            content: Base64-encoded DOCX content
            
        Returns:
            Extracted text
        """
        try:
            from docx import Document
            
            # Decode base64
            docx_bytes = base64.b64decode(content)
            docx_file = BytesIO(docx_bytes)
            
            # Extract text
            doc = Document(docx_file)
            text_parts = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)
            
            return "\n\n".join(text_parts)
            
        except Exception as e:
            return f"[DOCX parsing error: {str(e)}]"
